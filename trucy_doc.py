import os
import re
import zipfile
import xml.etree.ElementTree as ET


class TrucyDocLoader:
    """Absolute-path text/document loader for TXT, MD and DOCX files."""

    SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx"}

    @classmethod
    def INPUT_TYPES(cls):
        return {
            "required": {
                "file_path": (
                    "STRING",
                    {
                        "default": "C:\\scripts\\prompt.txt",
                        "multiline": False,
                        "tooltip": "支持 Windows 复制文件地址格式，例如：\"D:\\项目\\脚本.docx\"",
                    },
                ),
                "encoding": (
                    ["auto", "utf-8", "gb18030", "gbk", "utf-16"],
                    {"default": "auto"},
                ),
            }
        }

    RETURN_TYPES = ("STRING", "STRING", "STRING", "STRING", "STRING")
    RETURN_NAMES = ("text", "filename", "absolute_path", "extension", "modified_time")
    FUNCTION = "load_document"
    CATEGORY = "TrucyNodes/Text"

    @classmethod
    def IS_CHANGED(cls, file_path, encoding="auto"):
        path = cls._clean_path(file_path)

        try:
            stat = os.stat(path)
            # 同时使用路径、文件大小和纳秒级修改时间，避免文件修改后命中旧缓存。
            return (path, stat.st_size, stat.st_mtime_ns, encoding)
        except OSError:
            return (path, "FILE_NOT_FOUND", encoding)

    @staticmethod
    def _clean_path(file_path):
        """清理 Windows 复制文件地址产生的引号和空白。"""
        path = "" if file_path is None else str(file_path).strip()

        # 处理普通英文双引号："D:\\project\\test.txt"
        path = path.strip('"')

        # 兼容用户可能复制到的中文弯引号：“D:\\project\\test.txt”
        path = path.strip('“”')

        # 有些程序会在路径前后附加单引号，也一并兼容。
        path = path.strip("'")

        return os.path.expandvars(os.path.expanduser(path.strip()))

    @staticmethod
    def _read_text_file(path, encoding):
        if encoding != "auto":
            with open(path, "r", encoding=encoding, newline="") as file:
                return file.read()

        # utf-8-sig 可以自动移除 UTF-8 BOM。
        encodings = ["utf-8-sig", "gb18030", "gbk", "utf-16"]
        last_error = None

        for current_encoding in encodings:
            try:
                with open(path, "r", encoding=current_encoding, newline="") as file:
                    return file.read()
            except UnicodeDecodeError as error:
                last_error = error

        # 最后的容错读取，保证工作流不会因为少量异常字符直接中断。
        try:
            with open(path, "r", encoding="utf-8", errors="replace", newline="") as file:
                return file.read()
        except Exception:
            if last_error is not None:
                raise last_error
            raise

    @staticmethod
    def _read_docx_file(path):
        """
        读取 DOCX 中的正文和表格。
        优先使用 python-docx；如果环境没有安装，则使用 DOCX 内部 XML
        进行基础读取，避免节点完全不可用。
        """
        try:
            from docx import Document

            document = Document(path)
            parts = []

            for paragraph in document.paragraphs:
                text = paragraph.text
                if text.strip():
                    parts.append(text)

            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    parts.append(" | ".join(cells))

            return "\n".join(parts)

        except ImportError:
            # DOCX 本质上是 ZIP + XML。此回退方案可读取基础段落，
            # 但不会完整保留 Word 的高级排版信息。
            with zipfile.ZipFile(path, "r") as archive:
                xml_data = archive.read("word/document.xml")

            root = ET.fromstring(xml_data)
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            paragraphs = []

            for paragraph in root.iter(namespace + "p"):
                texts = []
                for node in paragraph.iter(namespace + "t"):
                    texts.append(node.text or "")
                paragraph_text = "".join(texts).strip()
                if paragraph_text:
                    paragraphs.append(paragraph_text)

            return "\n".join(paragraphs)

    @staticmethod
    def _normalize_text(text):
        # 统一 Windows、Linux 和 Mac 换行，去掉 UTF-8 BOM。
        text = str(text).replace("\ufeff", "")
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        return text.strip()

    def load_document(self, file_path, encoding):
        path = self._clean_path(file_path)

        if not path:
            error = "Error: file path is empty"
            return (error, "N/A", "", "", "")

        if not os.path.isfile(path):
            error = f"Error: file not found: {path}"
            return (error, "N/A", path, "", "")

        absolute_path = os.path.abspath(path)
        extension = os.path.splitext(absolute_path)[1].lower()
        filename = os.path.splitext(os.path.basename(absolute_path))[0]

        if extension not in self.SUPPORTED_EXTENSIONS:
            error = (
                "Error: unsupported file type. "
                "Only .txt, .md and .docx are supported."
            )
            return (error, filename, absolute_path, extension, "")

        try:
            if extension == ".docx":
                text = self._read_docx_file(absolute_path)
            else:
                text = self._read_text_file(absolute_path, encoding)

            text = self._normalize_text(text)
            modified_time = str(os.path.getmtime(absolute_path))

            print(
                f"[TrucyNodes] Document loaded: {absolute_path} "
                f"({len(text)} chars)"
            )

            return (
                text,
                filename,
                absolute_path,
                extension,
                modified_time,
            )

        except Exception as error:
            message = f"Error reading document: {error}"
            print(f"[TrucyNodes] {message}")
            return (message, filename, absolute_path, extension, "")


NODE_CLASS_MAPPINGS = {
    "TrucyDocLoader": TrucyDocLoader,
}

NODE_DISPLAY_NAME_MAPPINGS = {
    "TrucyDocLoader": "🚀 Document Loader (TXT / MD / DOCX) (Trucy)",
}